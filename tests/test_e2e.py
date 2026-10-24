"""Сквозные проверки на тестовом магазине: два снимка, сбои, robots.txt, отчёт."""
import contextlib
import io
import logging
import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

from mock_shop.shop import MockShop, expected_changes
from monitor.cli import load_config, main, run_crawl, run_report
from monitor.diff import compare, summarize
from monitor.fetch import Fetcher, FetchError
from monitor.store import Store

ROOT = pathlib.Path(__file__).resolve().parent.parent
logging.disable(logging.CRITICAL)


class E2E(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = str(pathlib.Path(self.tmp.name) / 'prices.db')
        self.shop = MockShop(day=1)
        self.base = self.shop.start()
        self.cfg = load_config(str(ROOT / 'config.json'), self.base)
        self.cfg['backoff_seconds'] = 0

    def tearDown(self):
        self.shop.stop()
        self.tmp.cleanup()

    def two_snapshots(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        self.shop.day = 2
        run_crawl(self.cfg, self.db, no_delay=True)

    def runs(self):
        st = Store(self.db)
        try:
            return st.db.execute('SELECT * FROM runs ORDER BY id').fetchall()
        finally:
            st.close()

    # ---- основной сценарий
    def test_snapshots_and_diff_match_the_shop(self):
        self.two_snapshots()
        r1, r2 = self.runs()
        self.assertEqual((r1['status'], r1['items'], r1['pages'], r1['duplicates']), ('ok', 36, 7, 4))
        self.assertEqual((r2['status'], r2['items'], r2['pages'], r2['duplicates']), ('ok', 37, 7, 7))
        st = Store(self.db)
        try:
            ok = st.last_ok_runs(2)
            changes = compare(st.load_items(ok[1]['id']), st.load_items(ok[0]['id']))
        finally:
            st.close()
        self.assertEqual(dict(summarize(changes)), {k: v for k, v in expected_changes().items() if v})

    def test_duplicates_are_taken_once_from_the_main_section(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        st = Store(self.db)
        try:
            items = st.load_items(1)
        finally:
            st.close()
        self.assertEqual(items['NB-1007'].category, 'Ноутбуки')            # не «Акции»
        self.assertEqual(items['NB-1007'].old_price_kop, 3299000)

    # ---- сбои
    def test_temporary_503_is_retried(self):
        self.shop.fail_once['/catalog/phones/?page=2'] = 2
        run_crawl(self.cfg, self.db, no_delay=True)
        self.assertEqual(self.runs()[0]['status'], 'ok')
        self.assertEqual(self.runs()[0]['items'], 36)
        self.assertEqual(self.shop.hits.count('/catalog/phones/?page=2'), 3)    # 2 отказа + успех

    def test_persistent_failure_makes_snapshot_partial_and_it_is_not_compared(self):
        self.two_snapshots()
        self.shop.fail_once['/catalog/phones/'] = 99
        run_crawl(self.cfg, self.db, no_delay=True)
        r = self.runs()
        self.assertEqual([x['status'] for x in r], ['ok', 'ok', 'partial'])
        self.assertIn('Смартфоны', r[2]['errors'])
        st = Store(self.db)
        try:
            self.assertEqual([x['id'] for x in st.last_ok_runs(2)], [2, 1])
        finally:
            st.close()

    def test_empty_catalog_is_flagged(self):
        self.shop.day = 99                                                   # магазин отдаёт пустые страницы
        run_crawl(self.cfg, self.db, no_delay=True)
        r = self.runs()[0]
        self.assertEqual(r['status'], 'partial')
        self.assertIn('каталог пуст', r['errors'])

    def test_sharp_drop_in_items_is_flagged(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        self.cfg['categories'] = {'Наушники': '/catalog/headphones/'}        # «пропали» три четверти каталога
        run_crawl(self.cfg, self.db, no_delay=True)
        r = self.runs()
        self.assertEqual([x['status'] for x in r], ['ok', 'partial'])
        self.assertIn('снимок помечен неполным', r[1]['errors'])

    # ---- robots.txt и пауза
    def test_robots_txt_is_respected(self):
        f = Fetcher(self.base, delay=0)
        self.assertFalse(f.allowed(self.base + '/admin/'))
        self.assertTrue(f.allowed(self.base + '/catalog/laptops/'))
        with self.assertRaises(FetchError):
            f.get(self.base + '/admin/')
        self.assertNotIn('/admin/', self.shop.hits)

    def test_crawl_delay_from_robots(self):
        self.assertEqual(Fetcher(self.base, delay=0.1).delay, 1.0)           # robots просит 1 с
        self.assertEqual(Fetcher(self.base, delay=0.1, honor_crawl_delay=False).delay, 0.1)

    def test_user_agent_is_declared_and_not_spoofed(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        self.assertGreater(len(self.shop.agents), 7)                         # robots.txt + все страницы
        self.assertEqual(set(self.shop.agents), {'PriceMonitor/1.0'})     # везде свой честный User-Agent

    # ---- отчёт
    def test_excel_report(self):
        self.two_snapshots()
        out = str(pathlib.Path(self.tmp.name) / 'r.xlsx')
        digest = run_report(self.cfg, self.db, out)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ['Сводка', 'Изменения', 'Прайс (снимок 2)'])
        total = sum(expected_changes().values())
        self.assertEqual(wb['Изменения'].max_row, total + 1)
        self.assertEqual(wb['Прайс (снимок 2)'].max_row, 37 + 1)
        summary = {r[0].value: r[1].value for r in wb['Сводка'].iter_rows(min_row=9, max_row=14, max_col=2)}
        self.assertEqual(summary['Цена снизилась'], 11)
        self.assertEqual(summary['Цена выросла'], 3)
        self.assertEqual(summary['Новый товар'], 3)
        self.assertEqual(len(wb['Сводка']._charts), 1)
        for ws in wb:                                                        # в отчёте нет ошибок и формул-пустышек
            for row in ws.iter_rows():
                for c in row:
                    self.assertFalse(isinstance(c.value, str) and c.value.startswith('#'), (ws.title, c.coordinate))
        self.assertIn('Подешевело (11)', digest)
        self.assertIn('Новых: 3', digest)

    def test_report_needs_two_good_snapshots(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        self.assertIsNone(run_report(self.cfg, self.db, str(pathlib.Path(self.tmp.name) / 'x.xlsx')))

    def test_cli_exit_codes(self):
        cfg_path, out = str(ROOT / 'config.json'), str(pathlib.Path(self.tmp.name) / 'cli.xlsx')
        args = ['--config', cfg_path, '--base-url', self.base, '--db', self.db, '--out', out, '--no-delay']
        with contextlib.redirect_stdout(io.StringIO()):        # CLI печатает сводку, в тестах она не нужна
            self.assertEqual(main(['report'] + args), 2)       # нечего сравнивать
            self.assertEqual(main(['run'] + args), 0)          # первый снимок: это нормально
            self.assertEqual(main(['run'] + args), 0)          # второй: отчёт без изменений
            self.assertTrue(pathlib.Path(out).exists())
            self.shop.fail_once['/catalog/laptops/'] = 99
            self.assertEqual(main(['crawl'] + args), 1)        # неполный снимок: код 1 для cron

    def test_report_without_changes(self):
        run_crawl(self.cfg, self.db, no_delay=True)
        run_crawl(self.cfg, self.db, no_delay=True)
        out = str(pathlib.Path(self.tmp.name) / 'same.xlsx')
        digest = run_report(self.cfg, self.db, out)
        wb = load_workbook(out)
        self.assertEqual(wb['Изменения'].max_row, 1)           # только заголовок
        self.assertEqual(len(wb['Сводка']._charts), 0)         # графика без данных нет
        self.assertIn('Новых: 0', digest)


if __name__ == '__main__':
    unittest.main()
