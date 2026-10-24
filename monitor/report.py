"""Отчёт в Excel: сводка с графиком, список изменений, полный прайс."""
from __future__ import annotations

from typing import Dict, List, Optional
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .diff import summarize
from .models import KIND_ORDER, KIND_RU, Change, Item

INK = '1F2A33'
HEAD_FILL = PatternFill('solid', fgColor='1F3A4D')
ZEBRA = PatternFill('solid', fgColor='F5F7F9')
BOTTOM = Border(bottom=Side(style='thin', color='D5DBE0'))
KIND_STYLE = {                       # цвет заливки, цвет текста
    'price_down': ('DDF1E4', '17623A'),
    'price_up': ('FBE0DC', '9E2A1D'),
    'new': ('DFEBFA', '1F4E8C'),
    'removed': ('E8E8E8', '555555'),
    'back_in_stock': ('D8F0EE', '14675F'),
    'out_of_stock': ('FCEBD2', '9A5B0B'),
    'price_set': ('EFEFEF', '555555'),
    'price_unset': ('EFEFEF', '555555'),
}
PRICE_KINDS = ('price_down', 'price_up', 'price_set', 'price_unset')
PCT_FMT = '+0.0%;-0.0%;0.0%'
LOCAL_HOSTS = {'127.0.0.1', 'localhost', '::1'}
CENTER = Alignment(horizontal='center', vertical='center')
RIGHT = Alignment(horizontal='right')


def F(**kw) -> Font:
    """Шрифт с явным именем: так отчёт одинаково выглядит в Excel, LibreOffice и просмотрщиках."""
    return Font(name='Calibri', size=kw.pop('size', 11), **kw)


def _rub(kop: Optional[int]):
    if kop is None:
        return None
    return kop // 100 if kop % 100 == 0 else kop / 100


def _money(cell, kop: Optional[int]) -> None:
    if kop is None:
        cell.value, cell.alignment = '—', RIGHT
        return
    cell.value = _rub(kop)
    cell.number_format = '#,##0.00 "₽"' if kop % 100 else '#,##0 "₽"'


def _dash(cell) -> None:
    cell.value, cell.alignment = '—', RIGHT


def _link(cell, url: str) -> None:
    """Локальный стенд: ссылки были бы мёртвыми, показываем путь текстом. Реальный сайт: кликабельная ссылка."""
    if not url:
        return
    u = urlsplit(url)
    if u.hostname in LOCAL_HOSTS:
        cell.value, cell.font = u.path, F(color='6B7783')
    else:
        cell.value, cell.hyperlink = url, url
        cell.font = F(color='0563C1', underline='single')


def _header(ws, row: int, titles: List[str]) -> None:
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font = HEAD_FILL, F(bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 26


def _widths(ws, widths: List[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _landscape(ws) -> None:
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_report(path: str, shop: str, prev_run: Dict, curr_run: Dict,
                 prev: Dict[str, Item], curr: Dict[str, Item], changes: List[Change]) -> None:
    wb = Workbook()
    _summary(wb.active, shop, prev_run, curr_run, changes)
    _changes_sheet(wb.create_sheet('Изменения'), changes)
    _price_sheet(wb.create_sheet('Прайс (снимок 2)'), curr)
    wb.save(path)


# ---------------------------------------------------------------- сводка
def _summary(ws, shop: str, prev_run, curr_run, changes: List[Change]) -> None:
    ws.title = 'Сводка'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '1F3A4D'
    ws['A1'] = 'Мониторинг цен: сравнение двух снимков'
    ws['A1'].font = F(size=16, bold=True, color=INK)
    ws['A2'] = (f'Магазин: {shop}')
    ws['A2'].font = F(italic=True, color='9A5B0B')
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 32

    _header(ws, 4, ['Снимок', 'Время начала', 'Товаров', 'Страниц', 'Предупреждений'])
    for r, (name, run) in enumerate((('№1 (предыдущий)', prev_run), ('№2 (текущий)', curr_run)), 5):
        values = [name, run['started_at'].replace('T', ' '), run['items'], run['pages'], run['warnings']]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = F(), BOTTOM
            cell.alignment = Alignment(horizontal='left' if c < 3 else 'center')

    _header(ws, 8, ['Что изменилось', 'Количество'])
    cnt = summarize(changes)
    r = 9
    for kind in KIND_ORDER:
        if not cnt[kind] and kind in ('price_set', 'price_unset'):
            continue
        fill, color = KIND_STYLE[kind]
        a = ws.cell(r, 1, KIND_RU[kind])
        a.fill, a.font = PatternFill('solid', fgColor=fill), F(bold=True, color=color)
        b = ws.cell(r, 2, cnt[kind])
        b.font, b.alignment = F(bold=True, color=INK), Alignment(horizontal='center')
        a.border = b.border = BOTTOM
        r += 1

    # таблица цен: от самого сильного снижения до самого сильного роста; из неё же строится график
    pcs = sorted((c for c in changes if c.kind in ('price_down', 'price_up')), key=lambda c: c.pct or 0)
    top = r + 2
    ws.cell(top - 1, 1, 'Изменения цен').font = F(bold=True, size=12, color=INK)
    _header(ws, top, ['Товар', 'Артикул', 'Было', 'Стало', 'Снижение', 'Рост'])
    for i, c in enumerate(pcs, top + 1):
        ws.cell(i, 1, c.title if len(c.title) <= 52 else c.title[:51].rstrip() + '…').font = F()
        ws.cell(i, 2, c.sku).font = F()
        _money(ws.cell(i, 3), c.old_kop)
        _money(ws.cell(i, 4), c.new_kop)
        for col in (3, 4):
            ws.cell(i, col).font = F()
        cell = ws.cell(i, 5 if c.kind == 'price_down' else 6, (c.pct or 0) / 100)
        cell.number_format, cell.font = PCT_FMT, F(bold=True, color=KIND_STYLE[c.kind][1])
        for col in range(1, 7):
            ws.cell(i, col).border = BOTTOM
    last = top + len(pcs)

    if pcs:
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap, ch.gapWidth = 'bar', 'clustered', 100, 40
        ch.title = 'Изменение цены, %'
        ch.style = 10
        ch.height, ch.width = 9.5 + 0.55 * len(pcs), 17
        ch.add_data(Reference(ws, min_col=5, max_col=6, min_row=top, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=2, min_row=top + 1, max_row=last))
        ch.series[0].graphicalProperties.solidFill = '2E8B57'
        ch.series[1].graphicalProperties.solidFill = 'C8442F'
        for s in ch.series:
            s.invertIfNegative = False
        ch.x_axis.scaling.orientation = 'maxMin'          # первая строка таблицы сверху
        ch.x_axis.tickLblPos = 'low'                      # подписи артикулов слева, а не поверх столбцов
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        ch.y_axis.number_format = '0%'
        ch.legend.position = 'b'
        ws.add_chart(ch, 'H4')

    _widths(ws, [52, 19, 14, 14, 17, 12, 3])
    _landscape(ws)


# ---------------------------------------------------------------- изменения
def _changes_sheet(ws, changes: List[Change]) -> None:
    ws.sheet_properties.tabColor = '2E8B57'
    _header(ws, 1, ['Что изменилось', 'Артикул', 'Товар', 'Раздел', 'Было', 'Стало',
                    'Изменение, ₽', 'Изменение, %', 'Страница товара'])
    for i, c in enumerate(changes, 2):
        fill, color = KIND_STYLE[c.kind]
        k = ws.cell(i, 1, KIND_RU[c.kind])
        k.fill, k.font = PatternFill('solid', fgColor=fill), F(bold=True, color=color)
        for col, v in ((2, c.sku), (3, c.title), (4, c.category)):
            ws.cell(i, col, v).font = F()
        _money(ws.cell(i, 5), c.old_kop)
        _money(ws.cell(i, 6), c.new_kop)
        d, p = ws.cell(i, 7), ws.cell(i, 8)
        if c.kind in PRICE_KINDS and c.delta_kop is not None:       # у смены наличия цена не меняется: прочерк
            d.value, d.number_format, d.font = _rub(c.delta_kop), '+#,##0 "₽";-#,##0 "₽"', F(color=color, bold=True)
            p.value, p.number_format, p.font = (c.pct or 0) / 100, PCT_FMT, F(color=color, bold=True)
        else:
            _dash(d)
            _dash(p)
        for col in (5, 6):
            ws.cell(i, col).font = F()
        _link(ws.cell(i, 9), c.url)
        for col in range(1, 10):
            ws.cell(i, col).border = BOTTOM
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:I{max(len(changes) + 1, 2)}'
    _widths(ws, [20, 11, 60, 13, 14, 14, 17, 17, 34])
    _landscape(ws)


# ---------------------------------------------------------------- прайс
def _price_sheet(ws, items: Dict[str, Item]) -> None:
    ws.sheet_properties.tabColor = '6B7783'
    _header(ws, 1, ['Артикул', 'Товар', 'Раздел', 'Цена', 'Старая цена', 'Скидка', 'Наличие', 'Примечание', 'Страница товара'])
    for i, it in enumerate(sorted(items.values(), key=lambda x: x.sku), 2):
        for col, v in ((1, it.sku), (2, it.title), (3, it.category)):
            ws.cell(i, col, v).font = F()
        _money(ws.cell(i, 4), it.price_kop)
        ws.cell(i, 4).font = F()
        if it.old_price_kop:
            _money(ws.cell(i, 5), it.old_price_kop)
            ws.cell(i, 5).font = F(color='6B7783')
        if it.price_kop and it.old_price_kop:
            d = ws.cell(i, 6, -(1 - it.price_kop / it.old_price_kop))
            d.number_format, d.font, d.alignment = '0%', F(color='17623A', bold=True), CENTER
        st = ws.cell(i, 7, 'в наличии' if it.in_stock else 'нет')
        st.font, st.alignment = F(color='17623A' if it.in_stock else '9A5B0B'), CENTER
        ws.cell(i, 8, it.note).font = F(italic=True, color='6B7783')
        _link(ws.cell(i, 9), it.url)
        for col in range(1, 10):
            ws.cell(i, col).border = BOTTOM
            if i % 2 == 1:
                ws.cell(i, col).fill = ZEBRA
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:I{max(len(items) + 1, 2)}'
    _widths(ws, [11, 60, 13, 14, 14, 9, 12, 20, 34])
    _landscape(ws)
